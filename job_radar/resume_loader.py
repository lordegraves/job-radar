from pathlib import Path

from docx import Document
from pypdf import PdfReader

from job_radar.config import ConfigError
from job_radar.normalize import clean_text


# Keep resume format support centralized here so CLI scans, resume matching,
# and future GUI uploads all normalize resumes through the same path.
SUPPORTED_RESUME_EXTENSIONS = {".docx", ".md", ".pdf", ".txt"}


def load_resume_text(path: str | Path) -> str:
    resume_path = Path(path)
    extension = resume_path.suffix.lower()

    if extension not in SUPPORTED_RESUME_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_RESUME_EXTENSIONS))
        raise ConfigError(
            f"Unsupported resume format: {extension or 'none'}. "
            f"Supported formats: {supported}"
        )

    if not resume_path.exists():
        raise ConfigError(f"Resume file does not exist: {resume_path}")

    text = _read_resume_text(resume_path, extension)
    normalized_text = clean_text(text)

    # A PDF can be a valid file but still produce no useful text, for example
    # if it is scanned as images. Treat that the same as an empty resume.
    if not normalized_text:
        raise ConfigError(f"Resume file is empty or unreadable: {resume_path}")

    return normalized_text


def write_normalized_resume_text(
    source_path: str | Path,
    normalized_text_path: str | Path,
) -> Path:
    normalized_text = load_resume_text(source_path)
    output_path = Path(normalized_text_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(normalized_text + "\n", encoding="utf-8")
    return output_path


def _read_resume_text(resume_path: Path, extension: str) -> str:
    if extension in {".md", ".txt"}:
        return resume_path.read_text(encoding="utf-8")

    if extension == ".docx":
        return _read_docx_text(resume_path)

    if extension == ".pdf":
        return _read_pdf_text(resume_path)

    raise ConfigError(f"Unsupported resume format: {extension or 'none'}")


def _read_docx_text(resume_path: Path) -> str:
    try:
        document = Document(resume_path)
    except Exception as error:
        raise ConfigError(f"Could not read DOCX resume file: {resume_path}") from error

    parts: list[str] = []

    # Most resumes are paragraphs, but some templates use tables for layout.
    # Read both so formatted DOCX resumes do not lose important content.
    for paragraph in document.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return "\n".join(parts)


def _read_pdf_text(resume_path: Path) -> str:
    try:
        reader = PdfReader(resume_path)
    except Exception as error:
        raise ConfigError(f"Could not read PDF resume file: {resume_path}") from error

    parts: list[str] = []

    # Text-based PDFs usually extract cleanly. Scanned/image-only PDFs usually
    # extract nothing; load_resume_text handles that after normalization.
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception as error:
            raise ConfigError(f"Could not extract text from PDF resume file: {resume_path}") from error

        if page_text:
            parts.append(page_text)

    return "\n".join(parts)