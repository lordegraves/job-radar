import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook

from job_radar.candidate_profile import CandidateProfile
from job_radar.collectors.greenhouse import CollectorError
from job_radar.cli import (
    _import_history_records,
    build_parser,
    handle_bootstrap_user_data,
    handle_import_history,
    handle_scan,
    handle_tracker_add,
    handle_tracker_list,
    handle_tracker_update,
)
from job_radar.scan_service import _find_profile_avoid_matches
from job_radar.job_history import (
    EXPECTED_HEADERS,
    SIMPLIFIED_HEADERS,
    JobHistoryRecord,
)
from job_radar.storage import initialize_database, upsert_job_history_record
from job_radar.email_sender import EmailSendResult
from job_radar.models import JobPosting
from job_radar.normalize import make_canonical_key, make_content_hash
from job_radar.tracker.tracker_models import ApplicationRecord
from job_radar.tracker.tracker_storage import upsert_application


def read_report_snapshot(report_file: Path) -> dict:
    snapshot_file = report_file.with_suffix(".json")
    return json.loads(snapshot_file.read_text(encoding="utf-8"))


def find_snapshot_job(snapshot: dict, title: str) -> dict:
    section_names = (
        "top_matches",
        "review_needed",
        "tracked_applications",
        "new_jobs",
        "passed_not_recommended",
    )

    for section_name in section_names:
        for job in snapshot[section_name]:
            if job["title"] == title:
                return job

    raise AssertionError(f"Snapshot job not found: {title}")


def count_job_posting_rows(database_file: Path) -> int:
    with sqlite3.connect(database_file) as connection:
        cursor = connection.execute("SELECT COUNT(*) FROM job_postings")
        return int(cursor.fetchone()[0])


def count_scan_run_rows(database_file: Path) -> int:
    with sqlite3.connect(database_file) as connection:
        cursor = connection.execute("SELECT COUNT(*) FROM scan_runs")
        return int(cursor.fetchone()[0])


def fetch_latest_scan_run(database_file: Path) -> sqlite3.Row:
    with sqlite3.connect(database_file) as connection:
        connection.row_factory = sqlite3.Row
        row = connection.execute(
            """
            SELECT *
            FROM scan_runs
            ORDER BY id DESC
            LIMIT 1
            """
        ).fetchone()

    assert row is not None
    return row


def fetch_scan_errors(database_file: Path) -> list[sqlite3.Row]:
    with sqlite3.connect(database_file) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """
            SELECT *
            FROM scan_errors
            ORDER BY id
            """
        ).fetchall()

    return rows


def count_job_history_rows(database_file: Path) -> int:
    with sqlite3.connect(database_file) as connection:
        cursor = connection.execute("SELECT COUNT(*) FROM job_history")
        return int(cursor.fetchone()[0])


def count_application_tracker_rows(database_file: Path) -> int:
    with sqlite3.connect(database_file) as connection:
        cursor = connection.execute("SELECT COUNT(*) FROM application_tracker")
        return int(cursor.fetchone()[0])


def fetch_job_history_status_outcomes(
    database_file: Path,
) -> list[tuple[str | None, str | None]]:
    with sqlite3.connect(database_file) as connection:
        rows = connection.execute(
            """
            SELECT status, outcome_category
            FROM job_history
            ORDER BY status, outcome_category
            """
        ).fetchall()

    return [(row[0], row[1]) for row in rows]


def fetch_application_tracker_status_outcomes(
    database_file: Path,
) -> list[tuple[str, str | None]]:
    with sqlite3.connect(database_file) as connection:
        rows = connection.execute(
            """
            SELECT status, outcome
            FROM application_tracker
            ORDER BY status, outcome
            """
        ).fetchall()

    return [(row[0], row[1]) for row in rows]


def count_cross_table_company_role_duplicates(database_file: Path) -> int:
    with sqlite3.connect(database_file) as connection:
        cursor = connection.execute(
            """
            SELECT COUNT(*)
            FROM job_history h
            JOIN application_tracker a
              ON lower(h.company) = lower(a.company_name)
             AND lower(h.role) = lower(a.role_title)
            """
        )

        return int(cursor.fetchone()[0])


def fetch_job_history_status(
    database_file: Path,
    import_key: str,
) -> str | None:
    with sqlite3.connect(database_file) as connection:
        cursor = connection.execute(
            """
            SELECT status
            FROM job_history
            WHERE import_key = ?
            """,
            (import_key,),
        )
        row = cursor.fetchone()

    if row is None:
        return None

    return str(row[0])


def make_cli_history_record(
    import_key: str,
    technical_match: str,
    outcome_category: str,
    primary_blocker: str | None = None,
) -> JobHistoryRecord:
    return JobHistoryRecord(
        history_type="Pipeline",
        company="Example AI",
        role="Senior Infrastructure Engineer",
        source="LinkedIn",
        ats_platform="Greenhouse",
        work_arrangement="Remote",
        location="Remote",
        comp_range="$160k-$200k",
        event_date="2026-06-01",
        status="Rejected - No Interview",
        outcome_category=outcome_category,
        recruiter_contact="Unknown",
        technical_match=technical_match,
        hiring_probability="Low",
        skills_signals="Linux, HPC, Infrastructure",
        primary_blocker=primary_blocker,
        secondary_blocker=None,
        revisit="No",
        include_in_job_radar=True,
        import_key=import_key,
        notes="Form rejection.",
    )


def write_cli_history_workbook(workbook_path: Path) -> None:
    workbook = Workbook()
    pipeline_sheet = workbook.active
    pipeline_sheet.title = "Pipeline Import"
    reviewed_sheet = workbook.create_sheet("Reviewed Import")

    pipeline_sheet.append(EXPECTED_HEADERS)
    reviewed_sheet.append(EXPECTED_HEADERS)

    pipeline_sheet.append(
        [
            "Pipeline",
            "Example AI",
            "Senior Infrastructure Engineer",
            "LinkedIn",
            "Greenhouse",
            "Remote",
            "Remote",
            "$160k-$200k",
            "2026-06-01",
            "Rejected - No Interview",
            "No Interview",
            "Unknown",
            "Very Strong",
            "Low",
            "Linux, HPC, Infrastructure",
            "Generic Remote Competition",
            None,
            "No",
            "Yes",
            "pipeline:example-ai:senior-infrastructure-engineer",
            "Form rejection.",
        ]
    )

    reviewed_sheet.append(
        [
            "Reviewed",
            "SkipCo",
            "Frontend Engineer",
            "LinkedIn",
            "Greenhouse",
            "Remote",
            "Remote",
            "$150k-$180k",
            "2026-06-02",
            "Skipped",
            "Skipped",
            None,
            "Weak",
            "Very Low",
            "Frontend",
            "Role Family Mismatch",
            None,
            "No",
            "Yes",
            "reviewed:skipco:frontend-engineer",
            "Not an infrastructure role.",
        ]
    )

    workbook.save(workbook_path)


def write_cli_simplified_history_workbook(workbook_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Job Log"
    worksheet.append(SIMPLIFIED_HEADERS)
    worksheet.append(
        [
            "jr-example-ai-001",
            "2026-07-03",
            "Example AI",
            "Senior Infrastructure Engineer",
            "https://example.com/jobs/123",
            "Job Radar",
            "Applied",
            "Pending / In Progress",
            "Jane Recruiter",
            "Applied from simplified Job Log.",
            "Yes",
        ]
    )
    worksheet.append(
        [
            None,
            "2026-07-03",
            "ManualCo",
            "Principal SRE",
            "https://example.com/manual-lead",
            "LinkedIn",
            "Passed",
            "N/A",
            None,
            "Manual lead from LinkedIn.",
            "Yes",
        ]
    )
    workbook.save(workbook_path)


def test_handle_import_history_imports_workbook_rows(
    tmp_path: Path,
    capsys,
) -> None:
    settings_file = tmp_path / "settings.yaml"
    workbook_file = tmp_path / "job-history.xlsx"
    database_file = tmp_path / "job_radar.sqlite3"

    write_cli_history_workbook(workbook_file)

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    handle_import_history(
        workbook_path=str(workbook_file),
        settings_path=str(settings_file),
    )

    output = capsys.readouterr().out

    assert database_file.exists()
    assert count_job_history_rows(database_file) == 2
    assert count_application_tracker_rows(database_file) == 0

    assert "Application history import complete" in output
    assert f"Workbook: {workbook_file}" in output
    assert f"Database: {database_file}" in output
    assert "Rows read: 2" in output
    assert "Rows imported: 2" in output
    assert "Rows updated: 0" in output
    assert "Rows skipped: 0" in output
    assert "Tracker rows imported: 0" in output
    assert "Tracker rows updated: 0" in output
    assert "Tracker rows skipped: 2" in output


def test_import_history_records_partitions_tracker_and_history_without_duplication(
    tmp_path: Path,
) -> None:
    database_file = tmp_path / "job_radar.sqlite3"
    initialize_database(database_file)

    records = [
        JobHistoryRecord(
            history_type="Pipeline",
            company="ActiveCo",
            role="Senior SRE",
            source="LinkedIn",
            ats_platform=None,
            work_arrangement=None,
            location=None,
            comp_range=None,
            event_date="2026-07-01",
            status="Applied",
            outcome_category="Pending / In Progress",
            recruiter_contact=None,
            technical_match=None,
            hiring_probability=None,
            skills_signals=None,
            primary_blocker=None,
            secondary_blocker=None,
            revisit=None,
            include_in_job_radar=True,
            import_key="job-radar-id:jr-activeco-senior-sre",
            notes="Active application.",
            job_radar_id="jr-activeco-senior-sre",
            posting_url="https://example.com/activeco/senior-sre",
            lead_source="LinkedIn",
        ),
        JobHistoryRecord(
            history_type="Pipeline",
            company="RejectedCo",
            role="Platform Engineer",
            source="LinkedIn",
            ats_platform=None,
            work_arrangement=None,
            location=None,
            comp_range=None,
            event_date="2026-07-02",
            status="Applied",
            outcome_category="Rejected - No Interview",
            recruiter_contact=None,
            technical_match=None,
            hiring_probability=None,
            skills_signals=None,
            primary_blocker=None,
            secondary_blocker=None,
            revisit=None,
            include_in_job_radar=True,
            import_key="job-radar-id:jr-rejectedco-platform-engineer",
            notes="Rejected before interview.",
            job_radar_id="jr-rejectedco-platform-engineer",
            posting_url="https://example.com/rejectedco/platform-engineer",
            lead_source="LinkedIn",
        ),
        JobHistoryRecord(
            history_type="Pipeline",
            company="PassedCo",
            role="Linux Administrator",
            source="LinkedIn",
            ats_platform=None,
            work_arrangement=None,
            location=None,
            comp_range=None,
            event_date="2026-07-03",
            status="Passed",
            outcome_category="N/A",
            recruiter_contact=None,
            technical_match=None,
            hiring_probability=None,
            skills_signals=None,
            primary_blocker=None,
            secondary_blocker=None,
            revisit=None,
            include_in_job_radar=True,
            import_key="job-radar-id:jr-passedco-linux-administrator",
            notes="Compensation too low.",
            job_radar_id="jr-passedco-linux-administrator",
            posting_url="https://example.com/passedco/linux-administrator",
            lead_source="LinkedIn",
        ),
    ]

    result = _import_history_records(
        database_path=str(database_file),
        records=records,
    )

    assert result == (2, 0, 1, 0, 2)
    assert count_job_history_rows(database_file) == 2
    assert count_application_tracker_rows(database_file) == 1
    assert count_cross_table_company_role_duplicates(database_file) == 0

    assert fetch_job_history_status_outcomes(database_file) == [
        ("Applied", "Rejected - No Interview"),
        ("Passed", "N/A"),
    ]
    assert fetch_application_tracker_status_outcomes(database_file) == [
        ("Applied", "Pending / In Progress"),
    ]


def test_handle_import_history_imports_simplified_workbook_rows(
    tmp_path: Path,
    capsys,
) -> None:
    settings_file = tmp_path / "settings.yaml"
    workbook_file = tmp_path / "job-history.xlsx"
    database_file = tmp_path / "job_radar.sqlite3"

    write_cli_simplified_history_workbook(workbook_file)

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    handle_import_history(
        workbook_path=str(workbook_file),
        settings_path=str(settings_file),
    )

    output = capsys.readouterr().out

    assert database_file.exists()
    assert count_job_history_rows(database_file) == 1
    assert count_application_tracker_rows(database_file) == 1

    assert fetch_job_history_status(
        database_file,
        "job-radar-id:jr-example-ai-001",
    ) is None
    assert fetch_job_history_status(
        database_file,
        "posting-url:https://example.com/manual-lead",
    ) == "Passed"

    assert "Application history import complete" in output
    assert f"Workbook: {workbook_file}" in output
    assert f"Database: {database_file}" in output
    assert "Rows read: 2" in output
    assert "Rows imported: 1" in output
    assert "Rows updated: 0" in output
    assert "Rows skipped: 0" in output
    assert "Tracker rows imported: 1" in output
    assert "Tracker rows updated: 0" in output
    assert "Tracker rows skipped: 1" in output


def test_handle_scan_records_completed_lifecycle(
    tmp_path: Path,
    monkeypatch,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"
    snapshot_file = tmp_path / "today.json"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10
  linux: 10

negative_keywords:
  sales: -10

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key="example-ai:senior-infrastructure-engineer:remote",
        content_hash="hash-linux-infrastructure",
    )

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        lambda company_config: [fake_posting],
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    scan_row = fetch_latest_scan_run(database_file)

    assert scan_row["status"] == "completed"
    assert scan_row["current_stage"] == "completed"
    assert scan_row["failure_summary"] is None
    assert scan_row["report_status"] == "completed"
    assert scan_row["email_status"] == "not_requested"
    assert scan_row["companies_requested"] == 1
    assert scan_row["companies_scanned"] == 1
    assert scan_row["jobs_collected"] == 1
    assert scan_row["collector_errors"] == 0
    assert fetch_scan_errors(database_file) == []
    assert snapshot_file.is_file()
    assert '"schema_version": 1' in snapshot_file.read_text(encoding="utf-8")


def test_handle_scan_records_collector_warning(
    tmp_path: Path,
    monkeypatch,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10

negative_keywords: {}

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    def fail_collection(company_config):
        raise CollectorError("temporary collector failure")

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fail_collection,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    scan_row = fetch_latest_scan_run(database_file)
    error_rows = fetch_scan_errors(database_file)

    assert scan_row["status"] == "completed_with_warnings"
    assert scan_row["current_stage"] == "completed"
    assert scan_row["companies_scanned"] == 1
    assert scan_row["jobs_collected"] == 0
    assert scan_row["collector_errors"] == 1
    assert scan_row["report_status"] == "completed"
    assert len(error_rows) == 1
    assert error_rows[0]["company_key"] == "example_ai"
    assert error_rows[0]["source_type"] == "greenhouse"
    assert error_rows[0]["error_type"] == "collection_error"
    assert error_rows[0]["error_message"] == "temporary collector failure"


def test_handle_scan_records_failed_stage(
    tmp_path: Path,
    monkeypatch,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    def fail_configuration(scoring_path):
        raise RuntimeError("invalid scoring configuration")

    monkeypatch.setattr(
        "job_radar.scan_service.load_scoring_config",
        fail_configuration,
    )

    try:
        handle_scan(
            config_path=str(config_file),
            settings_path=str(settings_file),
            report_path=str(report_file),
            scoring_path=str(tmp_path / "scoring.yaml"),
        )
    except RuntimeError as error:
        assert str(error) == "invalid scoring configuration"
    else:
        raise AssertionError("Expected handle_scan to raise RuntimeError.")

    scan_row = fetch_latest_scan_run(database_file)
    error_rows = fetch_scan_errors(database_file)

    assert scan_row["status"] == "failed"
    assert scan_row["current_stage"] == "configuration"
    assert scan_row["failure_summary"] == "invalid scoring configuration"
    assert scan_row["finished_at"] is not None
    assert len(error_rows) == 1
    assert error_rows[0]["error_type"] == "configuration_failure"
    assert error_rows[0]["error_message"] == "invalid scoring configuration"


def test_handle_scan_records_report_generation_failure(
    tmp_path: Path,
    monkeypatch,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10

negative_keywords: {}

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key="example-ai:senior-infrastructure-engineer:remote",
        content_hash="hash-linux-infrastructure",
    )

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        lambda company_config: [fake_posting],
    )

    def fail_report_generation(report_path, report):
        raise RuntimeError("report writer failed")

    monkeypatch.setattr(
        "job_radar.scan_service.write_html_report",
        fail_report_generation,
    )

    try:
        handle_scan(
            config_path=str(config_file),
            settings_path=str(settings_file),
            report_path=str(report_file),
            scoring_path=str(scoring_file),
        )
    except RuntimeError as error:
        assert str(error) == "report writer failed"
    else:
        raise AssertionError("Expected report generation to fail.")

    scan_row = fetch_latest_scan_run(database_file)
    error_rows = fetch_scan_errors(database_file)

    assert scan_row["status"] == "failed"
    assert scan_row["current_stage"] == "report_generation"
    assert scan_row["failure_summary"] == "report writer failed"
    assert scan_row["report_status"] == "not_started"
    assert len(error_rows) == 1
    assert error_rows[0]["error_type"] == "report_generation_failure"
    assert error_rows[0]["error_message"] == "report writer failed"


def test_handle_scan_collects_stores_scores_and_reports_jobs(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"
    html_report_file = tmp_path / "today.html"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10
  linux: 10

negative_keywords:
  sales: -10

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    canonical_key = make_canonical_key(
        company_key="example_ai",
        title="Senior Infrastructure Engineer",
        location="Remote",
    )

    content_hash = make_content_hash(
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key=canonical_key,
        content_hash=content_hash,
    )

    def fake_collect_jobs_for_company(company_config):
        return [fake_posting]

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fake_collect_jobs_for_company,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    output = capsys.readouterr().out
    report_text = html_report_file.read_text(encoding="utf-8")

    assert database_file.exists()
    assert not report_file.exists()
    assert html_report_file.exists()
    assert report_file.with_suffix(".json").exists()
    assert count_job_posting_rows(database_file) == 1
    assert count_scan_run_rows(database_file) == 1

    assert "Scan requested" in output
    assert "Companies enabled: 1" in output
    assert "Jobs collected: 1" in output
    assert "Actionable jobs stored: 1" in output
    assert "Jobs not actionable: 0" in output
    assert "Jobs new: 1" in output
    assert "Jobs seen: 0" in output
    assert "Jobs changed: 0" in output
    assert "Collector errors: 0" in output
    assert "<h1>Job Radar Report</h1>" in report_text
    assert "<strong>Companies enabled:</strong> 1" in report_text
    assert "<strong>Jobs collected:</strong> 1" in report_text
    assert "<strong>Actionable jobs stored:</strong> 1" in report_text
    assert "<strong>Jobs not actionable:</strong> 0" in report_text
    assert "<strong>New jobs:</strong> 1" in report_text
    assert "<strong>Seen jobs:</strong> 0" in report_text
    assert "<strong>Changed jobs:</strong> 0" in report_text
    assert "<strong>Collector errors:</strong> 0" in report_text
    assert "<strong>Top match score threshold:</strong> 1" in report_text
    assert "<strong>Review-needed score threshold:</strong> 100" in report_text

    assert "<h2>Top Matches</h2>" in report_text
    assert "<h2>Passed / Not Recommended</h2>" in report_text
    assert "<h2>All Jobs</h2>" not in report_text
    assert "Senior Infrastructure Engineer" in report_text
    assert "https://boards.greenhouse.io/exampleai/jobs/123" in report_text

    assert "<strong>Score:</strong> 40" in report_text
    assert (
        "<strong>Why this matched:</strong> infrastructure, linux, remote"
        in report_text
    )
    assert "Score reasons:" not in report_text
    assert "+30 title:infrastructure" not in report_text
    assert "+10 body:linux" not in report_text
    assert "+0 location_allowed:remote" not in report_text
    assert "<strong>Work location fit:</strong> remote" in report_text
    assert "Location status:" not in report_text

    assert "<strong>Company:</strong> Example AI" in report_text
    assert "<strong>Source:</strong> greenhouse" in report_text
    assert "<strong>Location:</strong> Remote" in report_text
    assert "https://boards.greenhouse.io/exampleai/jobs/123" in report_text
    assert (
        "<strong>Canonical key:</strong> "
        "<code>example-ai:senior-infrastructure-engineer:remote</code>"
        in report_text
    )


def test_handle_scan_adds_history_context_to_report(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10
  linux: 10

negative_keywords:
  sales: -10

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    initialize_database(database_file)
    upsert_job_history_record(
        database_file,
        make_cli_history_record(
            import_key="pipeline:example-ai:senior-infrastructure-engineer",
            technical_match="Strong",
            outcome_category="No Interview",
            primary_blocker="Compensation",
        ),
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key="example-ai:senior-infrastructure-engineer:remote",
        content_hash="hash-linux-infrastructure",
    )

    def fake_collect_jobs_for_company(company_config):
        return [fake_posting]

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fake_collect_jobs_for_company,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    capsys.readouterr()
    snapshot = read_report_snapshot(report_file)
    snapshot_job = find_snapshot_job(
        snapshot,
        "Senior Infrastructure Engineer",
    )

    assert not report_file.exists()
    assert report_file.with_suffix(".html").exists()
    assert report_file.with_suffix(".json").exists()
    assert snapshot_job["history_context"] == (
        "Prior similar application at Example AI ended "
        "No Interview despite Strong technical match"
    )
    assert snapshot_job["history_risk"] == (
        "caution: prior_no_interview_despite_strong_match"
    )


def test_handle_scan_imports_configured_history_workbook_before_report(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    workbook_file = tmp_path / "job-history.xlsx"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    write_cli_history_workbook(workbook_file)

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}
job_history_workbook_path: {workbook_file}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10
  linux: 10

negative_keywords:
  sales: -10

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key="example-ai:senior-infrastructure-engineer:remote",
        content_hash="hash-linux-infrastructure",
    )

    def fake_collect_jobs_for_company(company_config):
        return [fake_posting]

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fake_collect_jobs_for_company,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    output = capsys.readouterr().out
    snapshot = read_report_snapshot(report_file)
    snapshot_job = find_snapshot_job(
        snapshot,
        "Senior Infrastructure Engineer",
    )

    assert count_job_history_rows(database_file) == 2
    assert count_application_tracker_rows(database_file) == 0
    assert fetch_job_history_status(
        database_file,
        "pipeline:example-ai:senior-infrastructure-engineer",
    ) == "Rejected - No Interview"

    assert "Application history import complete" in output
    assert f"Workbook: {workbook_file}" in output
    assert "Rows read: 2" in output
    assert "Rows imported: 2" in output
    assert "Rows updated: 0" in output
    assert "Rows skipped: 0" in output
    assert "Tracker rows imported: 0" in output
    assert "Tracker rows updated: 0" in output
    assert "Tracker rows skipped: 2" in output

    assert not report_file.exists()
    assert report_file.with_suffix(".html").exists()
    assert report_file.with_suffix(".json").exists()
    assert snapshot_job["history_context"] == (
        "Prior similar application at Example AI ended "
        "No Interview despite Very Strong technical match"
    )
    assert snapshot_job["history_risk"] == (
        "caution: prior_no_interview_despite_strong_match"
    )


def test_handle_scan_attaches_existing_tracker_record_to_report(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10
  linux: 10

negative_keywords:
  sales: -10

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key="example-ai:senior-infrastructure-engineer:remote",
        content_hash="hash-linux-infrastructure",
    )

    initialize_database(database_file)
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id=fake_posting.job_radar_id,
            company_name="Example AI",
            role_title="Senior Infrastructure Engineer",
            source_url="https://boards.greenhouse.io/exampleai/jobs/123",
            status="applied",
            follow_up_on="2099-07-10",
            outcome="interviewing",
            notes="Already applied through company site.",
        ),
    )

    def fake_collect_jobs_for_company(company_config):
        return [fake_posting]

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fake_collect_jobs_for_company,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    capsys.readouterr()
    snapshot = read_report_snapshot(report_file)
    tracked_jobs = snapshot["tracked_applications"]

    assert count_application_tracker_rows(database_file) == 1
    assert snapshot["summary"]["tracked_applications"] == 1
    assert len(tracked_jobs) == 1
    assert tracked_jobs[0]["title"] == "Senior Infrastructure Engineer"
    assert tracked_jobs[0]["recommended_action"] == "Track Status"
    assert tracked_jobs[0]["action_rationale"] == (
        "You already applied for this job. "
        "Track the existing application instead of applying again."
    )

    with sqlite3.connect(database_file) as connection:
        tracker_row = connection.execute(
            """
            SELECT status, follow_up_on, outcome, notes
            FROM application_tracker
            WHERE job_radar_id = ?
            """,
            (fake_posting.job_radar_id,),
        ).fetchone()

    assert tracker_row == (
        "applied",
        "2099-07-10",
        "interviewing",
        "Already applied through company site.",
    )


def test_handle_scan_matches_tracked_application_by_source_url_when_ids_differ(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    config_file.write_text(
        """
companies:
  - company_key: runpod
    name: RunPod
    source_type: ashby
    source_slug: runpod
    enabled: true
""",
        encoding="utf-8",
    )
    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )
    scoring_file.write_text(
        """
title_keywords:
  infrastructure: 100
  reliability: 100

body_keywords:
  linux: 100

location:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:reliability
""",
        encoding="utf-8",
    )

    fake_posting = JobPosting(
        company_key="runpod",
        company_name="RunPod",
        source_type="ashby",
        source_job_id="1d14340c-c9bd-4754-80f4-5c83980cd413",
        source_url="https://jobs.ashbyhq.com/RunPod/1d14340c-c9bd-4754-80f4-5c83980cd413",
        title="Site Reliability Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key="runpod:site-reliability-engineer:remote",
        content_hash="hash-runpod-sre",
    )

    initialize_database(database_file)
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr_manual_runpod_site_reliability_engineer_a614aac51c",
            company_name="RunPod",
            role_title="Site Reliability Engineer",
            source_url="https://jobs.ashbyhq.com/RunPod/1d14340c-c9bd-4754-80f4-5c83980cd413",
            status="Applied",
            follow_up_on="2026-07-25",
            applied_on="2026-07-11",
            last_activity_on="2026-07-11",
            outcome="Pending / In Progress",
            notes="Applied through company site.",
        ),
    )

    def fake_collect_jobs_for_company(company_config):
        return [fake_posting]

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fake_collect_jobs_for_company,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    capsys.readouterr()
    snapshot = read_report_snapshot(report_file)
    tracked_jobs = snapshot["tracked_applications"]

    assert snapshot["top_matches"] == []
    assert snapshot["summary"]["tracked_applications"] == 1
    assert len(tracked_jobs) == 1
    assert tracked_jobs[0]["title"] == "Site Reliability Engineer"
    assert tracked_jobs[0]["url"] == (
        "https://jobs.ashbyhq.com/RunPod/"
        "1d14340c-c9bd-4754-80f4-5c83980cd413"
    )
    assert tracked_jobs[0]["recommended_action"] == "Track Status"

    with sqlite3.connect(database_file) as connection:
        tracker_row = connection.execute(
            """
            SELECT status, follow_up_on, outcome
            FROM application_tracker
            WHERE source_url = ?
            """,
            (fake_posting.source_url,),
        ).fetchone()

    assert tracker_row == (
        "Applied",
        "2026-07-25",
        "Pending / In Progress",
    )


def test_handle_scan_omits_track_status_for_untracked_job(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10
  linux: 10

negative_keywords:
  sales: -10

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key="example-ai:senior-infrastructure-engineer:remote",
        content_hash="hash-linux-infrastructure",
    )

    def fake_collect_jobs_for_company(company_config):
        return [fake_posting]

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fake_collect_jobs_for_company,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    capsys.readouterr()
    snapshot = read_report_snapshot(report_file)

    assert count_application_tracker_rows(database_file) == 0
    assert snapshot["summary"]["tracked_applications"] == 0
    assert snapshot["tracked_applications"] == []

    snapshot_job = find_snapshot_job(
        snapshot,
        "Senior Infrastructure Engineer",
    )
    assert snapshot_job["recommended_action"] != "Track Status"


def test_handle_scan_warns_and_continues_when_configured_history_workbook_is_missing(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    workbook_file = tmp_path / "missing-job-history.xlsx"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}
job_history_workbook_path: {workbook_file}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10
  linux: 10

negative_keywords:
  sales: -10

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key="example-ai:senior-infrastructure-engineer:remote",
        content_hash="hash-linux-infrastructure",
    )

    def fake_collect_jobs_for_company(company_config):
        return [fake_posting]

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fake_collect_jobs_for_company,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
    )

    output = capsys.readouterr().out

    assert database_file.exists()
    assert not report_file.exists()
    assert report_file.with_suffix(".html").exists()
    assert report_file.with_suffix(".json").exists()
    assert count_job_history_rows(database_file) == 0
    assert count_scan_run_rows(database_file) == 1

    assert "Application history import skipped" in output
    assert f"Workbook: {workbook_file}" in output
    assert "Reason: workbook file does not exist; using existing database history" in output
    assert "Scan summary:" in output
    assert "Actionable jobs stored: 1" in output


def test_handle_scan_passes_html_report_attachment_to_email_sender(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    config_file = tmp_path / "companies.yaml"
    settings_file = tmp_path / "settings.yaml"
    scoring_file = tmp_path / "scoring.yaml"
    database_file = tmp_path / "job_radar.sqlite3"
    report_file = tmp_path / "today.md"
    html_report_file = tmp_path / "today.html"

    config_file.write_text(
        """
companies:
  - company_key: example_ai
    name: Example AI
    source_type: greenhouse
    source_slug: exampleai
    enabled: true
""",
        encoding="utf-8",
    )

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7

email:
  enabled: false
  sender: ""
  sender_name: "Job Radar"
  recipients: []
  smtp_host: ""
  smtp_port: 587
""",
        encoding="utf-8",
    )

    scoring_file.write_text(
        """
positive_keywords:
  infrastructure: 10
  linux: 10

negative_keywords:
  sales: -10

location_preferences:
  allowed:
    remote: 100
  conditional: {}
  skipped: {}

top_matches:
  min_score: 1
  excluded_title_keywords: []
  strong_signals:
    - title:infrastructure
""",
        encoding="utf-8",
    )

    canonical_key = make_canonical_key(
        company_key="example_ai",
        title="Senior Infrastructure Engineer",
        location="Remote",
    )

    content_hash = make_content_hash(
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
    )

    fake_posting = JobPosting(
        company_key="example_ai",
        company_name="Example AI",
        source_type="greenhouse",
        source_job_id="123",
        source_url="https://boards.greenhouse.io/exampleai/jobs/123",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure.",
        canonical_key=canonical_key,
        content_hash=content_hash,
    )

    captured_email_call = {}

    def fake_collect_jobs_for_company(company_config):
        return [fake_posting]

    def fake_send_email_report(
        email_settings,
        subject,
        body,
        html_body=None,
        attachment_path=None,
    ):
        captured_email_call["email_settings"] = email_settings
        captured_email_call["subject"] = subject
        captured_email_call["body"] = body
        captured_email_call["html_body"] = html_body
        captured_email_call["attachment_path"] = attachment_path

        return EmailSendResult(
            sent=True,
            message="Email sent",
        )

    monkeypatch.setattr(
        "job_radar.scan_service.collect_jobs_for_company",
        fake_collect_jobs_for_company,
    )
    monkeypatch.setattr(
        "job_radar.scan_service.send_email_report",
        fake_send_email_report,
    )

    handle_scan(
        config_path=str(config_file),
        settings_path=str(settings_file),
        report_path=str(report_file),
        scoring_path=str(scoring_file),
        send_email=True,
    )

    output = capsys.readouterr().out

    assert not report_file.exists()
    assert html_report_file.exists()
    assert report_file.with_suffix(".json").exists()
    assert captured_email_call["attachment_path"] == html_report_file
    assert captured_email_call["subject"].startswith("Job Radar Report - ")
    assert "Full report:" in captured_email_call["body"]
    assert "Attached as HTML file." in captured_email_call["body"]
    assert str(report_file) not in captured_email_call["body"]
    assert captured_email_call["html_body"] is not None
    assert "<h1>Job Radar Report</h1>" in captured_email_call["html_body"]
    assert "View posting" in captured_email_call["html_body"]
    assert f"HTML report written: {html_report_file}" in output
    assert "Email send result: Email sent" in output


def test_find_profile_avoid_matches_detects_plain_avoid_terms() -> None:
    profile = CandidateProfile(
        name="Clayton Graves",
        compensation_floor_usd=160000,
        preferred_base_usd=185000,
        resume=None,
        core_strengths=[],
        credible_adjacent=[],
        learning_or_gap=[],
        avoid=["frontend", "product management"],
    )

    posting = JobPosting(
        company_key="test",
        company_name="Test Company",
        source_type="greenhouse",
        source_url="https://example.com/job",
        title="Senior Frontend Engineer",
        location="Remote",
        description="Build product management tooling.",
    )

    assert _find_profile_avoid_matches(profile, posting) == [
        "frontend",
        "product management",
    ]


def test_find_profile_avoid_matches_detects_cleared_only_roles() -> None:
    profile = CandidateProfile(
        name="Clayton Graves",
        compensation_floor_usd=160000,
        preferred_base_usd=185000,
        resume=None,
        core_strengths=[],
        credible_adjacent=[],
        learning_or_gap=[],
        avoid=["cleared-only roles"],
    )

    posting = JobPosting(
        company_key="test",
        company_name="Test Company",
        source_type="greenhouse",
        source_url="https://example.com/job",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Active Top Secret clearance and polygraph required.",
    )

    assert _find_profile_avoid_matches(profile, posting) == ["cleared-only roles"]


def test_find_profile_avoid_matches_ignores_non_matching_terms() -> None:
    profile = CandidateProfile(
        name="Clayton Graves",
        compensation_floor_usd=160000,
        preferred_base_usd=185000,
        resume=None,
        core_strengths=[],
        credible_adjacent=[],
        learning_or_gap=[],
        avoid=["frontend", "customer success", "cleared-only roles"],
    )

    posting = JobPosting(
        company_key="test",
        company_name="Test Company",
        source_type="greenhouse",
        source_url="https://example.com/job",
        title="Senior Infrastructure Engineer",
        location="Remote",
        description="Build Linux infrastructure and cluster systems.",
    )

    assert _find_profile_avoid_matches(profile, posting) == []


def test_parser_accepts_bootstrap_user_data_command() -> None:
    parser = build_parser()

    args = parser.parse_args(
        [
            "bootstrap-user-data",
            "--source-settings",
            "config/settings.yaml",
            "--source-profiles",
            "profiles",
            "--source-database",
            "data/job_radar.sqlite3",
            "--destination",
            "user-data",
        ]
    )

    assert args.command == "bootstrap-user-data"
    assert args.source_settings == "config/settings.yaml"
    assert args.source_profiles == "profiles"
    assert args.source_database == "data/job_radar.sqlite3"
    assert args.destination == "user-data"


def test_parser_uses_bootstrap_user_data_defaults() -> None:
    parser = build_parser()

    args = parser.parse_args(["bootstrap-user-data"])

    assert args.command == "bootstrap-user-data"
    assert args.source_settings == "config/settings.yaml"
    assert args.source_profiles == "profiles"
    assert args.source_database is None
    assert args.destination is None


def test_handle_bootstrap_user_data_copies_configuration(
    tmp_path: Path,
    capsys,
) -> None:
    source_root = tmp_path / "source"
    source_settings = source_root / "config" / "settings.yaml"
    source_profile = source_root / "profiles" / "clayton" / "profile.yaml"
    source_resume = source_root / "profiles" / "clayton" / "resume.md"
    source_database = source_root / "data" / "job_radar.sqlite3"
    destination = tmp_path / "user-data"
    source_settings.parent.mkdir(parents=True)
    source_profile.parent.mkdir(parents=True)
    source_database.parent.mkdir(parents=True)
    source_settings.write_text(
        "database_path: data/job_radar.sqlite3\n",
        encoding="utf-8",
    )
    source_profile.write_text(
        "candidate:\n  name: Clayton\n",
        encoding="utf-8",
    )
    source_resume.write_text("# Resume\n", encoding="utf-8")

    with sqlite3.connect(source_database) as connection:
        connection.execute(
            "CREATE TABLE bootstrap_marker (value TEXT NOT NULL)"
        )
        connection.execute(
            "INSERT INTO bootstrap_marker (value) VALUES (?)",
            ("copied through CLI",),
        )

    handle_bootstrap_user_data(
        source_settings_path=str(source_settings),
        source_profiles_path=str(source_root / "profiles"),
        source_database_path=str(source_database),
        destination=str(destination),
    )

    output = capsys.readouterr().out

    assert "User data bootstrap complete" in output
    assert f"Destination: {destination.resolve()}" in output
    assert "Files copied: 4" in output
    assert "Existing files preserved: 0" in output
    assert (destination / "config" / "settings.yaml").is_file()
    assert (destination / "profiles" / "clayton" / "profile.yaml").is_file()
    assert (destination / "profiles" / "clayton" / "resume.md").is_file()

    with sqlite3.connect(
        destination / "data" / "job_radar.sqlite3"
    ) as connection:
        marker_value = connection.execute(
            "SELECT value FROM bootstrap_marker"
        ).fetchone()[0]

    assert marker_value == "copied through CLI"


def test_handle_bootstrap_user_data_preserves_existing_files(
    tmp_path: Path,
    capsys,
) -> None:
    source_root = tmp_path / "source"
    source_settings = source_root / "config" / "settings.yaml"
    source_profile = source_root / "profiles" / "clayton" / "profile.yaml"
    destination = tmp_path / "user-data"
    destination_settings = destination / "config" / "settings.yaml"
    destination_profile = (
        destination / "profiles" / "clayton" / "profile.yaml"
    )
    source_settings.parent.mkdir(parents=True)
    source_profile.parent.mkdir(parents=True)
    destination_settings.parent.mkdir(parents=True)
    destination_profile.parent.mkdir(parents=True)
    source_settings.write_text("source settings\n", encoding="utf-8")
    source_profile.write_text("source profile\n", encoding="utf-8")
    destination_settings.write_text("existing settings\n", encoding="utf-8")
    destination_profile.write_text("existing profile\n", encoding="utf-8")

    handle_bootstrap_user_data(
        source_settings_path=str(source_settings),
        source_profiles_path=str(source_root / "profiles"),
        destination=str(destination),
    )

    output = capsys.readouterr().out

    assert "Files copied: 0" in output
    assert "Existing files preserved: 2" in output
    assert destination_settings.read_text(encoding="utf-8") == (
        "existing settings\n"
    )
    assert destination_profile.read_text(encoding="utf-8") == (
        "existing profile\n"
    )


def test_parser_accepts_grouped_history_import_command() -> None:
    parser = build_parser()

    args = parser.parse_args(
        [
            "history",
            "import",
            "--workbook",
            "data/job-history.xlsx",
            "--settings",
            "config/settings.yaml",
        ]
    )

    assert args.command == "history"
    assert args.history_command == "import"
    assert args.workbook == "data/job-history.xlsx"
    assert args.settings == "config/settings.yaml"


def test_parser_accepts_grouped_history_summary_command() -> None:
    parser = build_parser()

    args = parser.parse_args(
        [
            "history",
            "summary",
            "--settings",
            "config/settings.yaml",
        ]
    )

    assert args.command == "history"
    assert args.history_command == "summary"
    assert args.settings == "config/settings.yaml"


def test_parser_accepts_grouped_db_init_command() -> None:
    parser = build_parser()

    args = parser.parse_args(["db", "init"])

    assert args.command == "db"
    assert args.db_command == "init"


def test_parser_keeps_legacy_history_commands() -> None:
    parser = build_parser()

    import_args = parser.parse_args(
        [
            "import-history",
            "--workbook",
            "data/job-history.xlsx",
            "--settings",
            "config/settings.yaml",
        ]
    )
    summary_args = parser.parse_args(
        [
            "history-summary",
            "--settings",
            "config/settings.yaml",
        ]
    )
    init_args = parser.parse_args(["init-db"])

    assert import_args.command == "import-history"
    assert import_args.workbook == "data/job-history.xlsx"
    assert summary_args.command == "history-summary"
    assert init_args.command == "init-db"


def test_parser_accepts_tracker_list_command() -> None:
    parser = build_parser()

    args = parser.parse_args(
        [
            "tracker",
            "list",
            "--settings",
            "config/settings.yaml",
        ]
    )

    assert args.command == "tracker"
    assert args.tracker_command == "list"
    assert args.settings == "config/settings.yaml"
    assert args.needs_action is False
    assert args.needs_review is False


def test_parser_accepts_tracker_list_needs_action_command() -> None:
    parser = build_parser()

    args = parser.parse_args(
        [
            "tracker",
            "list",
            "--needs-action",
            "--settings",
            "config/settings.yaml",
        ]
    )

    assert args.command == "tracker"
    assert args.tracker_command == "list"
    assert args.needs_action is True
    assert args.needs_review is False
    assert args.settings == "config/settings.yaml"


def test_parser_accepts_tracker_list_needs_review_command() -> None:
    parser = build_parser()

    args = parser.parse_args(
        [
            "tracker",
            "list",
            "--needs-review",
            "--settings",
            "config/settings.yaml",
        ]
    )

    assert args.command == "tracker"
    assert args.tracker_command == "list"
    assert args.needs_action is False
    assert args.needs_review is True
    assert args.settings == "config/settings.yaml"


def test_handle_tracker_list_outputs_tracked_applications(
    tmp_path: Path,
    capsys,
) -> None:
    settings_file = tmp_path / "settings.yaml"
    database_file = tmp_path / "job_radar.sqlite3"

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    initialize_database(database_file)
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-stack-av-12345678",
            company_name="Stack AV",
            role_title="Senior Site Reliability Engineer",
            source_url="https://example.com/jobs/stack-av-sre",
            status="applied",
            follow_up_on="2099-07-10",
            outcome=None,
            notes="Applied through company site.",
        ),
    )

    handle_tracker_list(settings_path=str(settings_file))

    output = capsys.readouterr().out

    assert "Application tracker" in output
    assert f"Database: {database_file}" in output
    assert "Applications tracked: 1" in output
    assert "- Stack AV — Senior Site Reliability Engineer" in output
    assert "Job Radar ID: jr-stack-av-12345678" in output
    assert "Status: applied" in output
    assert "Workflow: follow_up_scheduled" in output
    assert "Follow up on: 2099-07-10" in output
    assert "URL: https://example.com/jobs/stack-av-sre" in output
    assert "Notes: Applied through company site." in output


def test_parser_accepts_tracker_update_command() -> None:
    parser = build_parser()

    args = parser.parse_args(
        [
            "tracker",
            "update",
            "--job-radar-id",
            "jr-stack-av-12345678",
            "--status",
            "applied",
            "--follow-up-on",
            "2099-07-10",
            "--applied-on",
            "2026-07-03",
            "--last-activity-on",
            "2026-07-05",
            "--outcome",
            "interviewing",
            "--notes",
            "Recruiter replied.",
            "--settings",
            "config/settings.yaml",
        ]
    )

    assert args.command == "tracker"
    assert args.tracker_command == "update"
    assert args.job_radar_id == "jr-stack-av-12345678"
    assert args.status == "applied"
    assert args.follow_up_on == "2099-07-10"
    assert args.applied_on == "2026-07-03"
    assert args.last_activity_on == "2026-07-05"
    assert args.outcome == "interviewing"
    assert args.notes == "Recruiter replied."
    assert args.settings == "config/settings.yaml"


def test_handle_tracker_list_filters_to_applications_needing_action(
    tmp_path: Path,
    capsys,
) -> None:
    settings_file = tmp_path / "settings.yaml"
    database_file = tmp_path / "job_radar.sqlite3"

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    initialize_database(database_file)
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-due-12345678",
            company_name="DueCo",
            role_title="Senior SRE",
            source_url="https://example.com/jobs/due",
            status="follow_up_due",
            follow_up_on="2026-07-04",
        ),
    )
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-active-12345678",
            company_name="ActiveCo",
            role_title="Infrastructure Engineer",
            source_url="https://example.com/jobs/active",
            status="interviewing",
        ),
    )
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-waiting-12345678",
            company_name="WaitingCo",
            role_title="Platform Engineer",
            source_url="https://example.com/jobs/waiting",
            status="applied",
        ),
    )

    handle_tracker_list(
        settings_path=str(settings_file),
        needs_action=True,
    )

    output = capsys.readouterr().out

    assert "Application tracker" in output
    assert "Filter: needs action" in output
    assert "Applications tracked: 2" in output
    assert "DueCo" in output
    assert "Workflow: follow_up_due" in output
    assert "ActiveCo" in output
    assert "Workflow: active_pipeline" in output
    assert "WaitingCo" not in output
    assert "Workflow: waiting" not in output


def test_handle_tracker_list_filters_to_applications_needing_review(
    tmp_path: Path,
    capsys,
) -> None:
    settings_file = tmp_path / "settings.yaml"
    database_file = tmp_path / "job_radar.sqlite3"

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    initialize_database(database_file)
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-invalid-date-12345678",
            company_name="DateReviewCo",
            role_title="Senior SRE",
            source_url="https://example.com/jobs/date-review",
            status="applied",
            follow_up_on="not-a-date",
        ),
    )
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-dormant-12345678",
            company_name="DormantCo",
            role_title="Infrastructure Engineer",
            source_url="https://example.com/jobs/dormant",
            status="applied",
            last_activity_on="2026-05-01",
        ),
    )
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-stale-12345678",
            company_name="StaleCo",
            role_title="Platform Engineer",
            source_url="https://example.com/jobs/stale",
            status="applied",
            last_activity_on="2026-03-01",
        ),
    )
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-presumed-closed-12345678",
            company_name="PresumedClosedCo",
            role_title="Linux Engineer",
            source_url="https://example.com/jobs/presumed-closed",
            status="applied",
            last_activity_on="2026-01-01",
        ),
    )
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-waiting-12345678",
            company_name="WaitingCo",
            role_title="Cluster Engineer",
            source_url="https://example.com/jobs/waiting",
            status="applied",
            last_activity_on="2026-07-01",
        ),
    )
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-active-12345678",
            company_name="ActiveCo",
            role_title="Infrastructure Engineer",
            source_url="https://example.com/jobs/active",
            status="interviewing",
        ),
    )

    handle_tracker_list(
        settings_path=str(settings_file),
        needs_review=True,
    )

    output = capsys.readouterr().out

    assert "Application tracker" in output
    assert "Filter: needs review" in output
    assert "Applications tracked: 4" in output
    assert "DateReviewCo" in output
    assert "Workflow: needs_date_review" in output
    assert "DormantCo" in output
    assert "Workflow: dormant" in output
    assert "StaleCo" in output
    assert "Workflow: stale" in output
    assert "PresumedClosedCo" in output
    assert "Workflow: presumed_closed" in output
    assert "WaitingCo" not in output
    assert "Workflow: waiting" not in output
    assert "ActiveCo" not in output
    assert "Workflow: active_pipeline" not in output


def test_handle_tracker_update_updates_existing_application(
    tmp_path: Path,
    capsys,
) -> None:
    settings_file = tmp_path / "settings.yaml"
    database_file = tmp_path / "job_radar.sqlite3"

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    initialize_database(database_file)
    upsert_application(
        database_file,
        ApplicationRecord(
            job_radar_id="jr-stack-av-12345678",
            company_name="Stack AV",
            role_title="Senior Site Reliability Engineer",
            source_url="https://example.com/jobs/stack-av-sre",
            status="review_needed",
            notes="Initial review.",
        ),
    )

    handle_tracker_update(
        settings_path=str(settings_file),
        job_radar_id="jr-stack-av-12345678",
        status="applied",
        follow_up_on="2099-07-10",
        applied_on="2026-07-03",
        last_activity_on="2026-07-05",
        outcome="interviewing",
        notes="Recruiter replied.",
    )

    output = capsys.readouterr().out

    assert "Application tracker updated" in output
    assert f"Database: {database_file}" in output
    assert "Job Radar ID: jr-stack-av-12345678" in output
    assert "Status: applied" in output
    assert "Follow up on: 2099-07-10" in output
    assert "Applied on: 2026-07-03" in output
    assert "Last activity on: 2026-07-05" in output
    assert "Outcome: interviewing" in output
    assert "Notes: Recruiter replied." in output

    handle_tracker_list(settings_path=str(settings_file))
    list_output = capsys.readouterr().out

    assert "Applications tracked: 1" in list_output
    assert "Status: applied" in list_output
    assert "Workflow: follow_up_scheduled" in list_output
    assert "Follow up on: 2099-07-10" in list_output
    assert "Applied on: 2026-07-03" in list_output
    assert "Last activity on: 2026-07-05" in list_output
    assert "Outcome: interviewing" in list_output
    assert "Notes: Recruiter replied." in list_output


def test_handle_tracker_update_reports_missing_application(
    tmp_path: Path,
    capsys,
) -> None:
    settings_file = tmp_path / "settings.yaml"
    database_file = tmp_path / "job_radar.sqlite3"

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    handle_tracker_update(
        settings_path=str(settings_file),
        job_radar_id="jr-missing-00000000",
        status="applied",
    )

    output = capsys.readouterr().out

    assert "Application tracker update failed" in output
    assert f"Database: {database_file}" in output
    assert "Job Radar ID: jr-missing-00000000" in output
    assert "Reason: tracked application was not found" in output


def test_parser_accepts_tracker_add_command() -> None:
    parser = build_parser()

    args = parser.parse_args(
        [
            "tracker",
            "add",
            "--job-radar-id",
            "jr-manual-12345678",
            "--company",
            "Example AI",
            "--role",
            "Senior Site Reliability Engineer",
            "--url",
            "https://example.com/jobs/123",
            "--status",
            "applied",
            "--follow-up-on",
            "2099-07-10",
            "--applied-on",
            "2026-07-03",
            "--last-activity-on",
            "2026-07-05",
            "--outcome",
            "interviewing",
            "--notes",
            "Applied through company site.",
            "--settings",
            "config/settings.yaml",
        ]
    )

    assert args.command == "tracker"
    assert args.tracker_command == "add"
    assert args.job_radar_id == "jr-manual-12345678"
    assert args.company == "Example AI"
    assert args.role == "Senior Site Reliability Engineer"
    assert args.url == "https://example.com/jobs/123"
    assert args.status == "applied"
    assert args.follow_up_on == "2099-07-10"
    assert args.applied_on == "2026-07-03"
    assert args.last_activity_on == "2026-07-05"
    assert args.outcome == "interviewing"
    assert args.notes == "Applied through company site."
    assert args.settings == "config/settings.yaml"


def test_handle_tracker_add_creates_tracked_application(
    tmp_path: Path,
    capsys,
) -> None:
    settings_file = tmp_path / "settings.yaml"
    database_file = tmp_path / "job_radar.sqlite3"

    settings_file.write_text(
        f"""
database_path: {database_file}
reports_path: {tmp_path}
logs_path: {tmp_path}

retention:
  report_retention_days: 90
  routine_event_retention_days: 90
  log_max_mb: 5
  log_backup_count: 5
  raw_capture_enabled: false
  raw_capture_retention_days: 7
""",
        encoding="utf-8",
    )

    handle_tracker_add(
        settings_path=str(settings_file),
        job_radar_id="jr-manual-12345678",
        company_name="Example AI",
        role_title="Senior Site Reliability Engineer",
        source_url="https://example.com/jobs/123",
        status="applied",
        follow_up_on="2099-07-10",
        applied_on="2026-07-03",
        last_activity_on="2026-07-05",
        outcome="interviewing",
        notes="Applied through company site.",
    )

    output = capsys.readouterr().out

    assert "Application tracker entry saved" in output
    assert f"Database: {database_file}" in output
    assert "Result: new" in output
    assert "Job Radar ID: jr-manual-12345678" in output
    assert "Company: Example AI" in output
    assert "Role: Senior Site Reliability Engineer" in output
    assert "Status: applied" in output
    assert "Follow up on: 2099-07-10" in output
    assert "Applied on: 2026-07-03" in output
    assert "Last activity on: 2026-07-05" in output
    assert "Outcome: interviewing" in output
    assert "URL: https://example.com/jobs/123" in output
    assert "Notes: Applied through company site." in output

    handle_tracker_list(settings_path=str(settings_file))
    list_output = capsys.readouterr().out

    assert "Applications tracked: 1" in list_output
    assert "- Example AI — Senior Site Reliability Engineer" in list_output
    assert "Job Radar ID: jr-manual-12345678" in list_output
    assert "Status: applied" in list_output
    assert "Applied on: 2026-07-03" in list_output
    assert "Last activity on: 2026-07-05" in list_output
