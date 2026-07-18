"""Tests supported resume formats, readable text extraction, and invalid files."""

import pytest
from docx import Document
from pypdf import PdfWriter

from job_radar.config import ConfigError
from job_radar.resume_loader import load_resume_text, write_normalized_resume_text


def test_load_resume_text_reads_markdown(tmp_path) -> None:
    resume_path = tmp_path / "resume.md"
    resume_path.write_text("# Resume\n\nLinux infrastructure and HPC operations", encoding="utf-8")

    assert load_resume_text(resume_path) == "# Resume Linux infrastructure and HPC operations"


def test_load_resume_text_reads_txt(tmp_path) -> None:
    resume_path = tmp_path / "resume.txt"
    resume_path.write_text("Cluster systems\nStorage operations", encoding="utf-8")

    assert load_resume_text(resume_path) == "Cluster systems Storage operations"


def test_load_resume_text_reads_docx(tmp_path) -> None:
    resume_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Linux infrastructure")
    document.add_paragraph("HPC operations")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cluster systems"
    table.cell(0, 1).text = "Storage operations"

    document.save(resume_path)

    assert load_resume_text(resume_path) == (
        "Linux infrastructure HPC operations Cluster systems Storage operations"
    )


def test_load_resume_text_rejects_empty_pdf(tmp_path) -> None:
    resume_path = tmp_path / "resume.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)

    with resume_path.open("wb") as output_file:
        writer.write(output_file)

    with pytest.raises(ConfigError, match="Resume file is empty or unreadable"):
        load_resume_text(resume_path)


def test_load_resume_text_rejects_unsupported_extension(tmp_path) -> None:
    resume_path = tmp_path / "resume.doc"
    resume_path.write_text("legacy word file", encoding="utf-8")

    with pytest.raises(ConfigError, match="Unsupported resume format: .doc"):
        load_resume_text(resume_path)


def test_load_resume_display_text_preserves_readable_lines(tmp_path) -> None:
    from job_radar.resume_loader import load_resume_display_text

    resume_path = tmp_path / "resume.md"
    resume_path.write_text(
        "# Resume\n\nLinux infrastructure\nHPC operations",
        encoding="utf-8",
    )

    assert load_resume_display_text(resume_path) == (
        "# Resume\n"
        "Linux infrastructure\n"
        "HPC operations"
    )


def test_write_normalized_resume_text(tmp_path) -> None:
    resume_path = tmp_path / "resume.md"
    normalized_path = tmp_path / "resume.normalized.txt"
    resume_path.write_text("Linux\n\nInfrastructure", encoding="utf-8")

    write_normalized_resume_text(resume_path, normalized_path)

    assert normalized_path.read_text(encoding="utf-8") == "Linux Infrastructure\n"
